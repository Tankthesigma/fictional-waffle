import base64
import builtins

import pandas as pd

from app.core.channel_inference import summarize_channels
from app.core.gating import rectangle_gate
from app.core.report_pdf import export_pdf_report
from app.core.report_pptx import export_pptx_report
from app.models.sample import SampleRecord
from app.ui.callbacks_reports import _export_report_figures, _report_status


def test_report_exports_create_files(tmp_path):
    frame = pd.DataFrame({"FSC-A": [1, 2], "SSC-A": [1, 2], "FL1-A": [10, 20]})
    sample = SampleRecord("s1", "s1.csv", path=tmp_path / "s1.csv", file_type="csv", events=frame)
    sample.channels = summarize_channels(frame)
    gate = rectangle_gate("g1", "main", "FSC-A", "SSC-A", 0, 3, 0, 3)
    stats = [{"gate_name": "main", "event_count": 2, "percent_total": 100.0}]
    image_path = tmp_path / "plot.png"
    image_path.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="))

    pdf = export_pdf_report([sample], [], [gate], stats, tmp_path / "report.pdf", figure_paths=[image_path])
    pptx = export_pptx_report([sample], [], [gate], stats, tmp_path / "report.pptx", figure_paths=[image_path])

    assert pdf.exists()
    assert pdf.stat().st_size > 0
    assert pptx.exists()
    assert pptx.stat().st_size > 0


def test_reports_include_comparison_rows(monkeypatch, tmp_path):
    frame = pd.DataFrame({"FSC-A": [1, 2], "SSC-A": [1, 2], "FL1-A": [10, 20]})
    sample = SampleRecord("s1", "s1.csv", path=tmp_path / "s1.csv", file_type="csv", events=frame)
    sample.channels = summarize_channels(frame)
    rows = [
        {
            "channel": "FL1-A",
            "control_median": 10.0,
            "treated_median": 20.0,
            "median_difference": 10.0,
            "fold_change": 2.0,
            "n_control": 2,
            "n_treated": 2,
            "notes": "exploratory only",
        }
    ]

    real_import = builtins.__import__

    def block_reportlab(name, *args, **kwargs):
        if name.startswith("reportlab"):
            raise ImportError("reportlab blocked for fallback test")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", block_reportlab)
    pdf = export_pdf_report([sample], [], [], [], tmp_path / "comparison.txt", comparison_rows=rows)
    monkeypatch.setattr(builtins, "__import__", real_import)
    pptx = export_pptx_report([sample], [], [], [], tmp_path / "comparison.pptx", comparison_rows=rows)

    pdf_text = pdf.read_text(encoding="utf-8")
    pptx_text = _pptx_text(pptx)
    assert "Exploratory Comparison" in pdf_text
    assert "FL1-A" in pdf_text
    assert "Exploratory Comparison" in pptx_text
    assert "median difference 10.0" in pptx_text


def test_plotly_static_image_renderer_available(tmp_path):
    import plotly.graph_objects as go

    path = tmp_path / "plot.png"
    fig = go.Figure(data=[go.Scatter(x=[1, 2, 3], y=[1, 4, 9])])

    fig.write_image(path, width=400, height=300, scale=1)

    assert path.exists()
    assert path.stat().st_size > 0


def test_report_figure_export_surfaces_static_image_errors(monkeypatch, tmp_path):
    frame = pd.DataFrame({"FSC-A": [1, 2, 3], "SSC-A": [1, 4, 9], "FL1-A": [10, 20, 30]})
    sample = SampleRecord("s1", "s1.csv", path=tmp_path / "s1.csv", file_type="csv", events=frame)
    sample.channels = summarize_channels(frame)

    def fail_write_image(self, *args, **kwargs):
        raise RuntimeError("renderer missing")

    monkeypatch.setattr("plotly.basedatatypes.BaseFigure.write_image", fail_write_image)

    result = _export_report_figures(
        "test-static-error",
        sample,
        [sample],
        "FSC-A",
        "SSC-A",
        "FL1-A",
        "scatter",
        "raw",
        150,
        1_000,
        [],
        False,
    )

    assert result.paths == []
    assert result.warnings
    status = _report_status("PDF", tmp_path / "report.pdf", result)
    assert "Static plot export needs review" in status
    assert "renderer missing" in status


def _pptx_text(path):
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                parts.append(shape.text)
    return "\n".join(parts)
