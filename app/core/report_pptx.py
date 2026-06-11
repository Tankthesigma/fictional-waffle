from __future__ import annotations

from datetime import datetime
from pathlib import Path

from app import __version__
from app.core.report_pdf import DISCLAIMER
from app.models.gate import GateDefinition
from app.models.qc_flag import QCFlag
from app.models.sample import SampleRecord


def export_pptx_report(
    samples: list[SampleRecord],
    qc_flags: list[QCFlag],
    gates: list[GateDefinition],
    gate_stats: list[dict[str, object]],
    output_path: str | Path,
    title: str = "Ask Flow Workbench Report",
    figure_paths: list[str | Path] | None = None,
) -> Path:
    """Export a PowerPoint report."""
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        from pptx import Presentation
    except Exception:
        target.write_text("python-pptx is not installed. Install requirements.txt to export PPTX.\n", encoding="utf-8")
        return target

    prs = Presentation()
    _title_slide(prs, title)
    _bullets_slide(prs, "Sample Summary", [f"{s.sample_id}: {s.event_count:,} events, {s.channel_count} channels" for s in samples])
    _bullets_slide(prs, "QC Summary", [f"{f.sample_id} {f.severity}: {f.title}" for f in qc_flags] or ["No QC flags currently present."])
    figures = _existing_figures(figure_paths)
    if figures:
        for figure_path in figures[:3]:
            _image_slide(prs, figure_path.stem.replace("-", " ").replace("_", " ").title(), figure_path)
    else:
        _bullets_slide(prs, "Representative Plots", ["No static plot images were available for this export."])
    _bullets_slide(prs, "Gate Statistics", [f"{row.get('gate_name')}: {row.get('event_count')} events" for row in gate_stats[:8]] or ["No gate statistics available."])
    _bullets_slide(prs, "Comparison", ["Control-versus-treated tables are exploratory and exported from the Compare tab when groups are selected."])
    _bullets_slide(prs, "Methods and Disclaimer", [f"App version {__version__}", DISCLAIMER])
    prs.save(target)
    return target


def _title_slide(prs, title: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = f"Generated {datetime.now().isoformat(timespec='seconds')}\nAsk Flow Workbench {__version__}"


def _bullets_slide(prs, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    for index, bullet in enumerate(bullets[:10]):
        para = body.paragraphs[0] if index == 0 else body.add_paragraph()
        para.text = bullet
        para.level = 0


def _image_slide(prs, title: str, image_path: Path) -> None:
    from pptx.util import Inches

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    slide.shapes.add_picture(str(image_path), Inches(0.7), Inches(1.25), width=Inches(8.9))


def _existing_figures(figure_paths: list[str | Path] | None) -> list[Path]:
    return [Path(path) for path in figure_paths or [] if Path(path).exists()]
